# Methods to read and verify Excel Lookahead.

# Lookahead headers:
# - Start Time
# - Well Name
# - Phase Code
# - Phase
# - Description
# - AFE Time
# - DSV Time
# - Actual Time

# Metadata:
# - Start Time (the very first start time of the lookahead)

import openpyxl
from settings import *

# Proposed workflow:
# - Identify latest lookahead. Ensure lookahead is a named table.
# - Use try-except to verify lookahead validity and raise errors.
# - Identify number of unique wells in the lookahead.
# - Compute Projection Time based on Actual Time, then AFE Time, then DSV Time.
# - Recalculate Projection Start Time based on Projection Time.
# - Generate Performance Tracker by well.
# - Generate well day fraction by date.

# Proposed verification:
# - Lookahead name, sheet name, and table name are as expected.
# - All expected Phase Codes and Phases for each well are present.
# - First phase of each well (i.e. phase 5) must have >1 day.
# - Raise warning if there are gaps in Actual Time.


def read_lookahead(excel_file_path=LATEST_LOOKAHEAD_DIR):
    sheet_name = 'Drilling Input'
    table_name = 'LookaheadTable'
    wb = openpyxl.load_workbook(filename=excel_file_path, data_only=True)
    ws = wb[sheet_name]
    lookup_table = ws.tables[table_name]
    data = ws[lookup_table.ref]
    rows_list = [[col.value for col in row] for row in data]
    df = pd.DataFrame(data=rows_list[1:], index=None, columns=rows_list[0])
    wb.close()
    return df


def generate_lookahead_projection(df):
    start_time = df['Start Time'].iloc[0]
    df['Projection Time'] = df['Actual Time'].fillna(df['AFE Time']).fillna(df['DSV Time'])
    df['Cumulative Projection Time'] = df['Projection Time'].cumsum().shift(fill_value=0)
    df['Projection Start Time'] = df['Cumulative Projection Time'].apply(lambda x: start_time + pd.Timedelta(hours=x))
    df['Projection End Time'] = df['Projection Start Time'] + pd.to_timedelta(df['Projection Time'], unit='hours')
    df.drop(['Cumulative Projection Time'], axis=1, inplace=True)
    # Round timestamp values to the nearest second.
    for col in df.select_dtypes(include='datetime64[ns]').columns:
        df[col] = df[col].dt.round('1s')
    return df


# Calculate intersection of two datetime ranges in number of days.
def calc_intersection(start1, end1, start2, end2):
    start = max(start1, start2)
    end = min(end1, end2)
    return (end - start).total_seconds() / (60 * 60 * 24) if start < end else 0


# Load workbook into dataframe.
df_lookahead = read_lookahead()

# Remove unnecessary columns.
df_lookahead = df_lookahead[['Start Time', 'Well Name', 'Phase Code', 'Phase', 'Description', 'AFE Time', 'DSV Time',
                             'Actual Time']]

# Identify number of unique wells in the lookahead.
lookahead_wells = {well for well in df_lookahead['Well Name'].unique() if well is not None}

# Recalculate projection start time.
df_lookahead = generate_lookahead_projection(df_lookahead)

# Identify well date range.
well_start_time = df_lookahead[df_lookahead['Phase Code'] > 0]['Projection Start Time'].iloc[0]
well_end_time = df_lookahead[df_lookahead['Phase Code'] > 0]['Projection End Time'].iloc[-1]
well_date_range = pd.date_range(start=well_start_time.date(), end=well_end_time.date())

# Generate performance tracker grouped by well.
grouped_df = df_lookahead.groupby(['Well Name', 'Phase Code', 'Phase']).agg(
    Projection_Start_Time=('Projection Start Time', 'first'),
    Projection_End_Time=('Projection End Time', 'last'),
    AFE_Time=('AFE Time', 'sum'),
    Actual_Time=('Actual Time', 'sum'))
grouped_df = grouped_df.reset_index()
grouped_df = grouped_df.rename({'Projection_Start_Time': 'Projection Start Time',
                                'Projection_End_Time': 'Projection End Time',
                                'AFE_Time': 'AFE Time',
                                'Actual_Time': 'Actual Time'}, axis='columns')

# Generate days ahead or behind.
grouped_df['Days Ahead/Behind'] = grouped_df.apply(
    lambda row: (row['Projection End Time'] - row['Projection Start Time']).total_seconds() / 86400 - row[
        'AFE Time'] / 24, axis=1)

# Merge WBS onto performance tracker.
grouped_df = grouped_df.merge(df_AFE_WBS.drop(['AFE Time'], axis=1), how='left')
grouped_df = grouped_df.sort_values(by=['Projection Start Time', 'Phase Code'])

# Generate day fraction per well phase.
for date in well_date_range:
    grouped_df[date.date()] = grouped_df.apply(
        lambda row: calc_intersection(row['Projection Start Time'], row['Projection End Time'], date,
                                      date + pd.Timedelta(days=1)), axis=1)
